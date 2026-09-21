"""Keep the strategy document aligned with the executable prototype."""
from pathlib import Path
import argparse
from docx.shared import RGBColor, Pt

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "strategy" / "Piter_Garcia_SaNDAI_8_Week_Seminar_Growth_Plan.docx"
parser = argparse.ArgumentParser()
parser.add_argument('--output', type=Path, default=SOURCE)
OUTPUT = parser.parse_args().output
SITE = "https://sandai-seminar-demo-qkpntto5ya-uc.a.run.app/"


def hyperlink(paragraph, text, url):
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.extend([properties, text_element])
    link.append(run)
    paragraph._p.append(link)


document = Document(SOURCE)

# Remove an earlier generated review block before rebuilding it. This keeps the
# script repeatable while preserving all author-written sections.
start = None
end = None
for index, paragraph in enumerate(document.paragraphs):
    if paragraph.text.strip() == "How to review the prototype":
        start = index
    if start is not None and paragraph.text.startswith("Access note:"):
        end = index
        break
if start is not None and end is not None:
    for paragraph in list(document.paragraphs[start:end + 1]):
        paragraph._element.getparent().remove(paragraph._element)

replacements = {
    "AI improves speed; Piter remains the accountable owner":
        "Proposed AI support roles; the prototype uses sample responses and deterministic analytics, not autonomous agents",
    "What the two notebooks prove": "What the prototype demonstrates",
    "Both implementations use small object-oriented components and visible checks so the logic remains testable, reusable, and understandable to a nontechnical reviewer.":
        "The website makes the two experiences accessible without setup. The supporting notebooks expose the reusable Python logic, inputs, and validation. The proposed operating workflow extends beyond the functions implemented in this prototype.",
    "Two interactive proofs-of-concept": "Two connected browser experiences",
    "1. Management Command Center - Piter_Garcia_SaNDAI_Management_Command_Center.ipynb": "1. Management Command Center",
    "2. Student Seminar Companion - Piter_Garcia_SaNDAI_Student_Seminar_Companion.ipynb": "2. Student Career Lab",
    "An executive-first command center with KPI cards and visual attendance, recovery, channel, learner-value, and conversion views. It uses synthetic data to demonstrate forecasting, decision gates, intervention, and documented ownership without changing the validated campaign mathematics.":
        "Review weekly attendance, cumulative pace, channels, and follow-on conversion. The sample contains 183 attendances but misses Week 3, so the full eight-week objective is not achieved. The Week 4 forecast and recovery case shows how I would act on risk; estimated improvement is not a guarantee.",
    "A UDL-aligned guided workshop with visual, quick-guide, and talk-then-write entry routes. Students reflect, practice, check an AI response, build truthful evidence, and leave with one realistic next action. The notebook requires no secret key and provides equivalent text for every required visual route.":
        "Work through Reflect, Practice, Check, Build, and Plan. Choose a community, career, or teaching scenario; review a sample AI response; explain a change; and download a takeaway with a realistic weekly plan. The website does not call a live AI service or certify completion.",
    "Uses synthetic data to demonstrate campaign forecasting, decision gates, channel/partner learning, ethical conversion, recovery, and a privacy-controlled learner-support layer.":
        "An executive-first command center with KPI cards and visual attendance, recovery, channel, learner-value, and conversion views. It uses synthetic data to demonstrate forecasting, decision gates, intervention, and documented ownership without changing the validated campaign mathematics.",
    "A shareable, no-secret-key learning tool for self-assessment, safe prompt practice, career/credential planning, a concise next-step plan, and optional voluntary feedback.":
        "A UDL-aligned guided workshop with visual, quick-guide, and talk-then-write entry routes. Students reflect, practice, check an AI response, build truthful evidence, and leave with one realistic next action. The notebook requires no secret key and provides equivalent text for every required visual route.",
}
for paragraph in document.paragraphs:
    if paragraph.text in replacements:
        paragraph.text = replacements[paragraph.text]

counting = 'Count each participant once per weekly seminar. Repeat attendance across weeks is allowed; assign one channel attribution per attendance. The prototype sums supplied channel totals and does not deduplicate individual check-ins.'
if not any(p.text == counting for p in document.paragraphs):
    next(p for p in document.paragraphs if p.text.strip() == 'Channel allocation rule').insert_paragraph_before(counting)

for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            if cell.text.startswith('Status logic\n'):
                cell.text = ('Forecast status\nRED: expected attendance below 20 or a supplied control flag is incomplete. '
                    'GREEN: expected attendance at least 22 and probability of 20 or more at least 80%. '
                    'YELLOW: otherwise. The demo assumes controls are complete; it does not inspect them. '
                    'T-7, T-5, and T-1 gates remain operator checks, not automated verification.')
            elif cell.text == 'No management data, login tracking, external transmission, or required disclosure; students choose what to share':
                cell.text = ('No management data or required personal disclosure. Web planning inputs are sent for calculation without application-level storage; '
                    'written reflections stay in the browser. The notebook does not call a live AI service.')
            elif cell.text == 'Notebook':
                cell.text = 'Experience and supporting notebook'

anchor = next(
    paragraph for paragraph in document.paragraphs
    if paragraph.text.strip() == "The weekly executive report stays concise"
)

heading = anchor.insert_paragraph_before("How to review the prototype", style="Heading 2")
heading.paragraph_format.page_break_before = True
intro = anchor.insert_paragraph_before(
    "Start with the website; no account, installation, or notebook execution is required. This plan explains the strategy and ownership. The browser experience makes the decisions and student value tangible. The repository and notebooks provide inspectable methods and sample inputs."
)

items = [
    ("1. Open the Management Command Center", SITE + "#management", "select Week 3 to see the miss, then review the Week 4 recovery case and channel contribution"),
    ("2. Try the Student Career Lab", SITE + "#student", "choose a scenario, complete the five review questions, then build and download a takeaway"),
    ("3. Review guide and supporting notebooks", "https://github.com/pzg8794/sandai-seminar-operations-demo/blob/main/START-HERE.md", "optional Colab notebooks and a reproducible route into the code"),
    ("4. Sample data and configuration", "https://github.com/pzg8794/sandai-seminar-operations-demo/tree/main/data/synthetic", "the illustrative inputs; configuration is linked from the review guide"),
    ("5. Requirement traceability", "https://github.com/pzg8794/sandai-seminar-operations-demo/blob/main/docs/REQUIREMENT_TRACEABILITY.md", "what the prototype demonstrates and what remains a proposed live workflow"),
]
for label, url, explanation in items:
    paragraph = anchor.insert_paragraph_before()
    hyperlink(paragraph, label, url)
    paragraph.add_run(f" — {explanation}.")

note = anchor.insert_paragraph_before(
    "Access note: the public website contains sample data, not a live student campaign. The private RIT Shared Drive supports local and Colab testing; reviewers do not need access to it. No enrollment, outreach, monetization, or student tracking is triggered by visiting the website."
)

# Preserve the original front page while adding a discoverable click target.
if not any('Open the interactive prototype' in p.text for p in document.paragraphs):
    front = next(p for p in document.paragraphs if p.text.strip() == 'The decision in one minute')
    link = front.insert_paragraph_before()
    hyperlink(link, 'Open the interactive prototype', SITE)
    link.add_run('  |  No sign-in required. Review steps and supporting evidence appear on page 7.')

# Match amended content to the retained design rather than losing run styles
# when replacing a paragraph or a table cell.
for paragraph in document.paragraphs:
    if paragraph.text in {'What the prototype demonstrates', 'Two connected browser experiences', 'How to review the prototype'}:
        for run in paragraph.runs:
            run.font.name = 'Cambria'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string('17365D')
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            if cell.text == 'Experience and supporting notebook':
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = 'Cambria'
                        r.font.size = Pt(9)
                        r.font.bold = True
                        r.font.color.rgb = RGBColor.from_string('FFFFFF')
            elif cell.text.startswith('No management data or required personal disclosure.'):
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = 'Cambria'
                        r.font.size = Pt(9)
document.save(OUTPUT)
print(OUTPUT)
